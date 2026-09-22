#!/usr/bin/env python3
"""Payroll journal-entry automation: raw Payworks CSV -> workpaper + DRAFT JE workbook.

Deterministic. No LLM, no network. stdlib + openpyxl only.

Usage:
    python payroll_automation.py <raw_payroll.csv> [--out-dir DIR]
        [--bank-leg-a 525.50] [--bank-leg-b 12160.00] [--bank-net-pay 14500.00]
        [--force]

Outputs (in --out-dir, default: current directory):
    PYROLL_REFERENCE_<period>.xlsx       reference workpaper, live formulas
    <period>_by_department_DRAFT.xlsx    QB-ready department JE draft

Validation checks (spec section 6) run before anything is written. If any fails
the script prints a diagnostic and exits 2 WITHOUT writing files, unless --force
is given (for a human who has reviewed the diagnostic and wants the files anyway).
Unknown department codes / unmapped labels are FLAGGED (highlighted + listed in
the NEEDS REVIEW notes), never guessed.
"""

import argparse
import csv
import sys
from collections import OrderedDict
from datetime import date
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------- #
# Fixed business logic (spec sections 2-5). Edit here, nowhere else.
# --------------------------------------------------------------------------- #

DEPT_NAMES = {
    '10': '0010-ACCOMODATIONS',
    '15': '0015-HOUSEKEEPING',
    '20': '0020-PINEWOODS',
    '25': '0025-DAYLODGE',            # seasonal
    '30': '0030-COUNTRY STORE',
    '40': '0040-BOAT HOUSE',
    '50': '0050-MANNING PARKS',
    '51': '0051-VISITOR CENTER',
    '52': '0052-SKYVIEW',
    '60': '0060-ALPINE',
    '70': '0070-NORDIC',              # seasonal
    '80': '0080-MAINTENANCE',
    '81': '0081-WEDDINGS & BANQUETS',
    '85': '0085-STAFF HOUSING',
    '90': '0090-GENERAL & ADMIN',
    '91': '0091-KITCHEN',
    '92': '0092-BEARS DEN',
    '93': '0093-MARKETING',
    '94': '0094-HUMAN RESOURCES',
    '95': '0095-ALPINE RENTALS',
    '96': '0096-ALPINE LESSONS',      # seasonal
    '97': '0097-INFORMATION TECHNOLOGY (IT)',
    '999': '0999-SKI PATROL',         # seasonal
    '0010-1': '0010-ACCOMODATIONS (Salaried)',
    '0020-1': '0020-PINEWOODS (Salaried)',
    '0050-1': '0050-MANNING PARKS (Salaried)',
    '0060-1': '0060-ALPINE (Salaried)',
}

SEASONAL_DEPTS = {'25', '70', '96', '999'}

# Department matrix: (account number, category name) -> source labels. ALL DEBIT.
GL_CATEGORY_SOURCE_LABELS = OrderedDict([
    ((6050, 'WAGES & SALARIES'), [
        'Regular Pay', 'Regular Salary', 'Wages', 'Overtime', 'Sick Pay', 'Training',
        'Double Time', 'Retro Regular Pay', 'Stat Pay @1.0', 'Stat Pay @1.5',
    ]),
    ((6046, 'CPP EXPENSE'), ['CPP/QPP Employer', 'CPP2/QPP2 Employer']),
    ((6047, 'EI EXPENSE'), ['EI Employer']),
    ((6049, 'VACATION EXPENSE'), ['Vac Each Pay']),
    ((6044, 'TIPS & GRATUITIES'), ['Tips & Gratitude']),
    ((2024, 'STAFF DEPOSIT'), ['Refund Staff Deposit']),
    ((6043, 'BENEFITS'), ['Benefits']),
    ((2011, 'VACATION PAYABLE'), []),   # always 0 - control row
])

# JSON key used in expected_results.json for each category, plus the QB account
# name used on the DRAFT workbook (must be an exact key of payroll_je.ACCOUNT_NAME_MAP).
CATEGORY_META = {
    6050: ('6050_wages_and_salaries', '6050 Wages & Salaries'),
    6046: ('6046_cpp_expense', '6046 CPP Expense'),
    6047: ('6047_ei_expense', '6047 EI Expense'),
    6049: ('6049_vacation_expense', '6049 Vacation expense'),
    6044: ('6044_tips_and_gratuities', '6044 Tips & Gratuities'),
    2024: ('2024_staff_deposit', '2024 Staff Deposits'),
    6043: ('6043_benefits', '6043 Benefits'),
    2011: ('2011_vacation_payable', '2011 Vacation Payable'),
}

LABEL_TO_CATEGORY = {
    label: acct for (acct, _name), labels in GL_CATEGORY_SOURCE_LABELS.items() for label in labels
}

# Top-block formulas. Each entry: (json key, QB account name, [source labels]).
# Labels that are absent in a period simply contribute nothing.
CPP_PAYABLE_LABELS = ['CPP/QPP Employee', 'CPP2/QPP2 Employee', 'CPP/QPP Employer', 'CPP2/QPP2 Employer']
EI_PAYABLE_LABELS = ['EI Employee', 'EI Employer']
FED_TAX_LABELS = ['Federal Tax']

LEG_A = [
    ('5004_dues_and_subscription', '5004 Dues and Subscription', ['Service Fees']),
    ('2031_gst_paid_on_purchases', '2031 GST Paid on Purchases', ['GST']),
    ('2040_pst_paid_on_purchases', '2040 PST Paid on Purchases', ['PST']),
]
LEG_B = [
    ('2013_ei_payable', '2013 EI Payable', EI_PAYABLE_LABELS),
    ('2014_cpp_payable', '2014 CPP Payable', CPP_PAYABLE_LABELS),
    ('2017_federal_income_tax_payable', '2017 Federal Income Tax Payable', FED_TAX_LABELS),
]
BANK_ACCOUNT = ('1004_envision_bank_account_credit', '1004 Envision Bank Account')

HEALTH_BENEFIT_LABELS = [
    'AD&D', 'Critical Illness', 'Dental', 'EAP', 'Extended Health', 'Life Insurance', 'LTD',
    'AD&D ER', 'Critical Illness ER', 'Dental ER', 'EAP ER', 'Extended Health ER',
    'Life Insurance ER', 'LTD ER',
]
WAGES_PAYABLE_LABELS = ['Net Pay', 'Prov. Garnishment', 'Net Pay Deduction', 'Excess Deductions']
UTILITIES_LABELS = ['Utilities 100', 'Utilities 170', 'Utilities Repayment']

CLOSING = [   # ALL CREDIT
    ('6007_utilities', '6007 Utilities', UTILITIES_LABELS),
    ('2024_staff_deposit_company', '2024 Staff Deposits', ['Staff Deposit']),
    ('2016_wages_payable', '2016 Wages Payable', WAGES_PAYABLE_LABELS),
    ('2017_federal_income_tax_payable', '2017 Federal Income Tax Payable', FED_TAX_LABELS),
    ('2014_cpp_payable', '2014 CPP Payable', CPP_PAYABLE_LABELS),
    ('2013_ei_payable', '2013 EI Payable', EI_PAYABLE_LABELS),
    ('2015_rrsp_payable', '2015 RRSP Payable', ['RRSP EE', 'RRSP ER']),
    ('6043_employer_health_benefits', '6043 Employer Health Benefits', HEALTH_BENEFIT_LABELS),
]

CENT = Decimal('0.01')
FILL_REVIEW = PatternFill('solid', fgColor='FFC7CE')
FILL_MANUAL = PatternFill('solid', fgColor='FFFF00')
BOLD = Font(bold=True)
UNMAPPED_MARK = '⚠️ UNMAPPED'


def D(x):
    return Decimal(str(x)).quantize(CENT, rounding=ROUND_HALF_UP)


# --------------------------------------------------------------------------- #
# Parsing (spec section 1)
# --------------------------------------------------------------------------- #

def normalize_dept_code(raw):
    code = raw.strip()
    if '-' in code:          # salaried codes like 0010-1: keep verbatim
        return code
    stripped = code.lstrip('0')
    return stripped if stripped else '0'


def parse_csv(path):
    """Returns (header_date, top_block, dept_blocks, dept_order, closing_rows_ignored).

    top_block: list of (batch, label, Decimal amount) in file order.
    dept_blocks: OrderedDict code -> list of (batch, label, Decimal amount).
    """
    with open(path, newline='', encoding='utf-8-sig') as f:
        rows = [r for r in csv.reader(f)]
    if not rows:
        raise ValueError('empty input file')
    header_date = rows[0][0].strip()
    if not (header_date.isdigit() and len(header_date) == 8):
        raise ValueError(f'row 1 must be a YYYYMMDD batch date, got {header_date!r}')

    top_block, dept_blocks, ignored = [], OrderedDict(), []
    mode = 'top'
    for i, r in enumerate(rows[1:], start=2):
        if not any(c.strip() for c in r):
            continue
        if len(r) < 5:
            raise ValueError(f'row {i}: expected Batch,DeptCode,,Label,Amount; got {r}')
        batch, dept, label, amount = r[0].strip(), r[1].strip(), r[3].strip(), r[4].strip()
        try:
            amt = D(amount)
        except Exception:
            raise ValueError(f'row {i}: non-numeric amount {amount!r}')
        if dept == '':
            if mode == 'top':
                top_block.append((batch, label, amt))
            else:
                mode = 'closing'
                ignored.append((batch, label, amt))
        else:
            mode = 'dept'
            code = normalize_dept_code(dept)
            dept_blocks.setdefault(code, []).append((batch, label, amt))
    return header_date, top_block, dept_blocks, list(dept_blocks.keys()), ignored


# --------------------------------------------------------------------------- #
# Computation (spec sections 2-5)
# --------------------------------------------------------------------------- #

def compute(header_date, top_block, dept_blocks, dept_order, ignored=None, bank=None):
    """Pure computation. Returns a dict with every figure plus flags/notes."""
    top = OrderedDict()             # label -> summed Decimal (first-seen order)
    for _b, label, amt in top_block:
        top[label] = top.get(label, Decimal('0')) + amt

    def tsum(labels):
        return sum((top[l] for l in labels if l in top), Decimal('0'))

    flags, notes = [], []
    dept_names = OrderedDict()
    matrix = OrderedDict()          # code -> OrderedDict(acct -> Decimal)
    unmapped = []                   # (code, label, amount)
    for code in dept_order:
        if code in DEPT_NAMES:
            dept_names[code] = DEPT_NAMES[code]
        else:
            dept_names[code] = f'DEPT {code} (name not confirmed)'
            flags.append({'type': 'unknown_department_code', 'department_code': code})
        cats = OrderedDict((acct, Decimal('0')) for (acct, _n) in GL_CATEGORY_SOURCE_LABELS)
        for _b, label, amt in dept_blocks[code]:
            acct = LABEL_TO_CATEGORY.get(label)
            if acct is None:
                unmapped.append((code, label, amt))
                flags.append({'type': 'unmapped_pay_component_label', 'label': label,
                              'department_code': code, 'amount': amt})
            else:
                cats[acct] += -amt      # dept lines are negative in the raw file; matrix is magnitude
        matrix[code] = cats

    matrix_total = sum((sum(c.values(), Decimal('0')) for c in matrix.values()), Decimal('0'))
    category_totals = OrderedDict(
        (acct, sum((m[acct] for m in matrix.values()), Decimal('0')))
        for (acct, _n) in GL_CATEGORY_SOURCE_LABELS)

    leg_a = OrderedDict((k, tsum(labels)) for k, _n, labels in LEG_A)
    leg_a[BANK_ACCOUNT[0]] = sum(leg_a.values(), Decimal('0'))
    leg_b = OrderedDict((k, tsum(labels)) for k, _n, labels in LEG_B)
    leg_b[BANK_ACCOUNT[0]] = sum(leg_b.values(), Decimal('0'))

    closing = OrderedDict((k, tsum(labels)) for k, _n, labels in CLOSING)
    closing_total = sum(closing.values(), Decimal('0'))

    cpp_employer_top = tsum(['CPP/QPP Employer', 'CPP2/QPP2 Employer'])
    ei_employer_top = tsum(['EI Employer'])

    # Period-specific quirks
    if any(l in top for l in ('CPP2/QPP2 Employee', 'CPP2/QPP2 Employer')):
        notes.append('CPP2/QPP2 (enhanced CPP) lines present this period; included in CPP Payable / CPP Expense.')
    present_hb = [l for l in HEALTH_BENEFIT_LABELS if l in top]
    if not present_hb:
        notes.append('Employer Health Benefits: no AD&D/Dental/etc. deduction lines in raw data - confirmed $0 for this period (checked, not skipped).')
    else:
        notes.append(f'Employer Health Benefits built from: {", ".join(present_hb)}.')
    present_wp = [l for l in WAGES_PAYABLE_LABELS if l in top]
    notes.append(f'Wages Payable built from: {", ".join(present_wp)}.')
    if 'Utilities Repayment' in top:
        notes.append('Utilities Repayment line present this period; included in 6007 Utilities.')
    seasonal = [c for c in dept_order if c in SEASONAL_DEPTS]
    if seasonal:
        notes.append('Seasonal departments present this period: ' + ', '.join(f'{c} ({DEPT_NAMES[c]})' for c in seasonal) + '.')

    # Validation (section 6)
    failures = []
    def check(name, left_name, left, right_name, right):
        if left != right:
            failures.append(f'{name}: {left_name} = {left}  vs  {right_name} = {right}  (diff {left - right})')
    check('CHECK 1 dept matrix total == closing total', 'matrix', matrix_total, 'closing', closing_total)
    check('CHECK 2 CPP Expense matrix == top-block CPP employer', 'matrix', category_totals[6046], 'top block', cpp_employer_top)
    check('CHECK 3 EI Expense matrix == top-block EI employer', 'matrix', category_totals[6047], 'top block', ei_employer_top)
    check('CHECK 4 Leg A debits == credit', 'debits', sum(list(leg_a.values())[:-1], Decimal('0')), 'credit', leg_a[BANK_ACCOUNT[0]])
    check('CHECK 5 Leg B debits == credit', 'debits', sum(list(leg_b.values())[:-1], Decimal('0')), 'credit', leg_b[BANK_ACCOUNT[0]])

    bank_status = {}
    net_pay = top.get('Net Pay', Decimal('0'))
    if bank:
        for key, ours, label in (('leg_a', leg_a[BANK_ACCOUNT[0]], 'Leg A (Dues/GST/PST)'),
                                 ('leg_b', leg_b[BANK_ACCOUNT[0]], 'Leg B (EI/CPP/FedTax)'),
                                 ('net_pay', net_pay, 'Net Pay')):
            supplied = bank.get(key)
            if supplied is None:
                bank_status[key] = 'not yet confirmed'
                continue
            if D(supplied) == ours:
                bank_status[key] = 'confirmed against bank statement'
            else:
                bank_status[key] = f'MISMATCH: bank {D(supplied)} vs computed {ours}'
                failures.append(f'CHECK 6 bank {label}: bank = {D(supplied)}  vs  computed = {ours}')
    else:
        bank_status = {k: 'not yet confirmed' for k in ('leg_a', 'leg_b', 'net_pay')}
    bank_confirmed = all(v.startswith('confirmed') for v in bank_status.values())
    notes.append('Bank confirmation: ' + ('all three withdrawals confirmed against bank statement.'
                                          if bank_confirmed else
                                          '; '.join(f'{k} {v}' for k, v in bank_status.items()) + '.'))

    return {
        'header_date': header_date,
        'top': top, 'top_block': top_block, 'ignored': ignored or [],
        'dept_order': dept_order, 'dept_names': dept_names, 'dept_blocks': dept_blocks,
        'matrix': matrix, 'category_totals': category_totals, 'matrix_total': matrix_total,
        'leg_a': leg_a, 'leg_b': leg_b,
        'closing': closing, 'closing_total': closing_total,
        'cpp_employer_top': cpp_employer_top, 'ei_employer_top': ei_employer_top,
        'net_pay': net_pay,
        'unmapped': unmapped, 'flags': flags, 'notes': notes, 'failures': failures,
        'bank_status': bank_status, 'bank_confirmed': bank_confirmed,
    }


# --------------------------------------------------------------------------- #
# Output 1 - reference workpaper (live formulas)
# --------------------------------------------------------------------------- #

def _ref_formula(rows, negate=False):
    if not rows:
        return '=0'
    inner = '+'.join(f'E{r}' for r in rows)
    return f'=-({inner})' if negate else f'={inner}'


def write_reference(res, path):
    """Reference workpaper laid out like the hand-built "Master Reference" workbook:
    raw source rows verbatim down columns A/B/D/E (batch, dept code, label, amount)
    with a per-department SUM subtotal in column F, a GL-coding box in columns
    H/K/L alongside the top block (Leg A, Leg B, Closing entries - all live
    formulas referencing the raw E-column cells), a department cross-tab matrix,
    and a reversal-check box that ties the trailing "closing noise" block back to
    the top-block Service Fees/GST/PST figures for a human to eyeball.
    """
    wb = Workbook()
    ws = wb.active
    period = res['header_date']
    try:
        d = date(int(period[:4]), int(period[4:6]), int(period[6:8]))
        ws.title = d.strftime('%B %d Payroll')[:31]
    except Exception:
        ws.title = 'Reference'
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['D'].width = 26
    ws.column_dimensions['E'].width = 13
    ws.column_dimensions['H'].width = 30

    ws['A1'] = int(period) if period.isdigit() else period
    ws['A1'].font = BOLD

    # ---- Top block: raw pass-through (A=batch, D=label, E=amount) ------------
    r = 2
    top_rows = {}                    # label -> [rows]
    for batch, label, amt in res['top_block']:
        ws.cell(r, 1, batch); ws.cell(r, 4, label); ws.cell(r, 5, float(amt))
        top_rows.setdefault(label, []).append(r)
        r += 1
    top_block_last = r - 1

    def top_ref(labels):
        rows = [rr for l in labels for rr in top_rows.get(l, [])]
        return _ref_formula(rows)

    # ---- GL coding box (H/K/L), alongside the top block, live formulas -------
    cr = 4
    cr += 1
    leg_a_first = cr
    for _k, name, labels in LEG_A:
        ws.cell(cr, 8, name); ws.cell(cr, 11, top_ref(labels)); cr += 1
    ws.cell(cr, 8, BANK_ACCOUNT[1]); ws.cell(cr, 12, f'=SUM(K{leg_a_first}:K{cr - 1})')
    leg_a_bank_row = cr
    cr += 2
    leg_b_first = cr
    for _k, name, labels in LEG_B:
        ws.cell(cr, 8, name); ws.cell(cr, 11, top_ref(labels)); cr += 1
    ws.cell(cr, 8, BANK_ACCOUNT[1]); ws.cell(cr, 12, f'=SUM(K{leg_b_first}:K{cr - 1})')
    leg_b_bank_row = cr
    cr += 2
    closing_first = cr
    closing_credit_rows = {}
    for key, name, labels in CLOSING:
        ws.cell(cr, 8, name)
        if key == '2014_cpp_payable':
            ws.cell(cr, 12, f'=K{leg_b_first + 1}')          # reuse Leg B's CPP calc
        elif key == '2013_ei_payable':
            ws.cell(cr, 12, f'=K{leg_b_first}')              # reuse Leg B's EI calc
        else:
            ws.cell(cr, 12, top_ref(labels))
        closing_credit_rows[key] = cr
        cr += 1
    ws.cell(cr, 8, 'CLOSING TOTAL').font = BOLD
    ws.cell(cr, 12, f'=SUM(L{closing_first}:L{cr - 1})').font = BOLD
    closing_total_cell = f'L{cr}'
    coding_box_last = cr

    r = max(r, coding_box_last) + 2

    # ---- Department blocks: raw pass-through + F-column subtotal -------------
    dept_item_rows = {}              # code -> {label: [rows]}
    dept_first_row, dept_last_row = {}, {}
    for code in res['dept_order']:
        name = res['dept_names'][code]
        if code not in DEPT_NAMES:
            hdr = ws.cell(r, 4, f'{code} — {name}'); hdr.font = BOLD; hdr.fill = FILL_REVIEW
            r += 1
        first = r
        dept_item_rows[code] = {}
        dept_code_val = code if '-' in code else (int(code) if code.isdigit() else code)
        for batch, label, amt in res['dept_blocks'][code]:
            ws.cell(r, 1, batch); ws.cell(r, 2, dept_code_val)
            ws.cell(r, 4, label); ws.cell(r, 5, float(amt))
            if label not in LABEL_TO_CATEGORY:
                ws.cell(r, 4).fill = FILL_REVIEW
                ws.cell(r, 6, UNMAPPED_MARK).fill = FILL_REVIEW
            dept_item_rows[code].setdefault(label, []).append(r)
            r += 1
        dept_first_row[code], dept_last_row[code] = first, r - 1
        ws.cell(r, 6, f'=SUM(E{first}:E{r - 1})' if r > first else '=0').font = BOLD
        r += 1

    # ---- Trailing reversal / closing-noise block (informational only) --------
    ignored_rows = {}
    if res['ignored']:
        r += 1
        ws.cell(r, 1, 'Below: trailing reversal block from the source export - bookkeeping '
                'noise, NOT part of the GL (see reversal check to the right).').font = BOLD
        r += 1
        for batch, label, amt in res['ignored']:
            ws.cell(r, 1, batch); ws.cell(r, 4, label); ws.cell(r, 5, float(amt))
            ignored_rows.setdefault(label, []).append(r)
            r += 1

    # ---- Department cross-tab matrix ------------------------------------------
    r += 2
    ws.cell(r, 1, 'DEPARTMENT MATRIX (magnitudes, all DEBIT)').font = BOLD; r += 1
    for j, code in enumerate(res['dept_order']):
        ws.cell(r, 2 + j, code).font = BOLD
    total_col = 2 + len(res['dept_order'])
    ws.cell(r, total_col, 'TOTAL').font = BOLD
    r += 1
    matrix_first = r
    for (acct, cname), labels in GL_CATEGORY_SOURCE_LABELS.items():
        ws.cell(r, 1, f'{acct} {cname}')
        for j, code in enumerate(res['dept_order']):
            rows = [rr for l in labels for rr in dept_item_rows[code].get(l, [])]
            ws.cell(r, 2 + j, _ref_formula(rows, negate=True))
        ws.cell(r, total_col, f'=SUM(B{r}:{get_column_letter(total_col - 1)}{r})')
        r += 1
    ws.cell(r, 1, 'GRAND TOTAL').font = BOLD
    for j in range(len(res['dept_order']) + 1):
        col = get_column_letter(2 + j)
        ws.cell(r, 2 + j, f'=SUM({col}{matrix_first}:{col}{r - 1})').font = BOLD
    matrix_total_cell = f'{get_column_letter(total_col)}{r}'
    r += 1
    ws.cell(r, 1, 'CHECK matrix - closing (must be 0)').font = BOLD
    ws.cell(r, 2, f'={matrix_total_cell}-{closing_total_cell}').font = BOLD
    r += 2

    # ---- Reversal check box: ties the ignored block back to the top block ----
    if res['ignored']:
        ws.cell(r, 8, 'CLOSING REVERSAL CHECK (must be 0)').font = BOLD; r += 1
        check_first = r
        for label in ('Service Fees', 'GST', 'PST'):
            if label in ignored_rows and label in top_rows:
                ws.cell(r, 8, label)
                ws.cell(r, 11, f'={_ref_formula(top_rows[label])[1:]}+{_ref_formula(ignored_rows[label])[1:]}')
                r += 1
        if 'Payroll Clearing Account' in ignored_rows:
            first_pca = top_rows.get('Payroll Clearing Account', [])
            last_pca = ignored_rows.get('Payroll Clearing Account', [])
            if first_pca and last_pca:
                ws.cell(r, 8, 'Payroll Clearing Account')
                ws.cell(r, 11, f'=E{first_pca[0]}+E{last_pca[0]}')
                r += 1
        ws.cell(r, 8, 'Total (must be $0.00)').font = BOLD
        ws.cell(r, 11, f'=SUM(K{check_first}:K{r - 1})').font = BOLD
        r += 2

    # ---- Bank rec box ----------------------------------------------------------
    ws.cell(r, 1, 'BANK REC').font = BOLD; r += 1
    ws.cell(r, 1, 'NetPay'); ws.cell(r, 2, top_ref(['Net Pay'])); net_row = r; r += 1
    ws.cell(r, 1, 'Bank (manual entry)')
    ws.cell(r, 2).fill = FILL_MANUAL
    ws.cell(r, 3, res['bank_status']['net_pay'])
    r += 1
    ws.cell(r, 1, 'Difference'); ws.cell(r, 2, f'=B{net_row}-B{r - 1}'); r += 2

    if res['flags']:
        ws.cell(r, 1, 'NEEDS REVIEW').font = BOLD; r += 1
        for line in review_lines(res):
            c = ws.cell(r, 1, line); c.fill = FILL_REVIEW; r += 1
    wb.save(path)


# --------------------------------------------------------------------------- #
# Output 2 - QB-ready department JE draft
# --------------------------------------------------------------------------- #

def review_lines(res):
    lines = []
    for f in res['flags']:
        if f['type'] == 'unknown_department_code':
            lines.append(f"Unknown department code {f['department_code']} -> shown as "
                         f"'DEPT {f['department_code']} (name not confirmed)'. Confirm the class name before posting.")
    for code, label, amt in res['unmapped']:
        lines.append(f"Unmapped pay component '{label}' in department {code}, amount {amt} "
                     f"-> {UNMAPPED_MARK}. Not included in any GL bucket; map it before posting.")
    return lines


def write_draft(res, path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'DRAFT JE'
    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 14
    ws.column_dimensions['D'].width = 60

    ws['A1'] = f"Payroll Journal Entry — batch {res['header_date']} — by department"
    ws['A1'].font = Font(bold=True, size=13)
    banner = ('DRAFT — built from raw payroll CSV, bank-confirmed (all three withdrawals matched).'
              if res['bank_confirmed'] else
              'DRAFT — built from raw payroll CSV, not a posted JE.')
    ws['A2'] = banner
    ws['A2'].font = Font(bold=True, color='C00000')
    r = 4
    subtotal_rows = []
    for code in res['dept_order']:
        name = res['dept_names'][code]
        hdr = ws.cell(r, 1, f'{code} — {name}'); hdr.font = BOLD
        if code not in DEPT_NAMES:
            hdr.fill = FILL_REVIEW
        r += 1
        for j, h in enumerate(('Account', 'Debit', 'Credit', 'Description'), start=1):
            ws.cell(r, j, h).font = BOLD
        r += 1
        first = r
        for (acct, cname), _labels in GL_CATEGORY_SOURCE_LABELS.items():
            amt = res['matrix'][code][acct]
            if amt == 0:
                continue
            ws.cell(r, 1, CATEGORY_META[acct][1]); ws.cell(r, 2, float(amt))
            ws.cell(r, 4, f'{name} — {cname.title()}')
            r += 1
        for ucode, label, amt in res['unmapped']:
            if ucode != code:
                continue
            c = ws.cell(r, 1, UNMAPPED_MARK); c.fill = FILL_REVIEW
            ws.cell(r, 2, float(-amt)).fill = FILL_REVIEW
            ws.cell(r, 4, f"UNMAPPED label '{label}' — not in GL_CATEGORY_SOURCE_LABELS, map before posting").fill = FILL_REVIEW
            r += 1
        ws.cell(r, 1, 'Subtotal').font = BOLD
        ws.cell(r, 2, f'=SUM(B{first}:B{r - 1})' if r > first else '=0').font = BOLD
        subtotal_rows.append(r)
        r += 2

    ws.cell(r, 1, 'Department grand total').font = BOLD
    dept_total_cell = f'B{r}'
    ws.cell(r, 2, '=' + ('+'.join(f'B{s}' for s in subtotal_rows) if subtotal_rows else '0')).font = BOLD
    r += 2

    ws.cell(r, 1, 'Company-wide closing entries').font = BOLD; r += 1
    for j, h in enumerate(('Account', 'Debit', 'Credit', 'Description'), start=1):
        ws.cell(r, j, h).font = BOLD
    r += 1
    first = r
    bank_desc = {
        '2016_wages_payable': 'net_pay', '2017_federal_income_tax_payable': 'leg_b',
        '2014_cpp_payable': 'leg_b', '2013_ei_payable': 'leg_b',
    }
    for key, name, _labels in CLOSING:
        amt = res['closing'][key]
        if amt == 0:
            desc = 'known-$0 period (no source lines in raw data; confirmed, not skipped)'
        elif key in bank_desc:
            desc = res['bank_status'][bank_desc[key]]
            if key == '2016_wages_payable' and desc.startswith('confirmed'):
                desc += ' (Net Pay component)'
        else:
            desc = 'internal-consistency figure only (not bank-matched)'
        ws.cell(r, 1, name); ws.cell(r, 3, float(amt)); ws.cell(r, 4, desc)
        r += 1
    ws.cell(r, 1, 'Closing total').font = BOLD
    ws.cell(r, 3, f'=SUM(C{first}:C{r - 1})').font = BOLD
    closing_total_cell = f'C{r}'
    r += 2

    ws.cell(r, 1, 'Combined grand total').font = BOLD
    ws.cell(r, 2, f'={dept_total_cell}').font = BOLD
    ws.cell(r, 3, f'={closing_total_cell}').font = BOLD
    r += 1
    ws.cell(r, 1, 'Check (Debit - Credit, must be $0.00)').font = BOLD
    ws.cell(r, 2, f'={dept_total_cell}-{closing_total_cell}').font = BOLD
    r += 2

    ws.cell(r, 1, 'NOTES — READ BEFORE POSTING').font = BOLD; r += 1
    n = 1
    for line in review_lines(res) + res['notes']:
        c = ws.cell(r, 1, f'{n}. {line}')
        if n <= len(review_lines(res)):
            c.fill = FILL_REVIEW
        r += 1; n += 1
    wb.save(path)


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #

def run(csv_path, out_dir='.', bank=None, force=False, quiet=False):
    header_date, top_block, dept_blocks, dept_order, ignored = parse_csv(csv_path)
    res = compute(header_date, top_block, dept_blocks, dept_order, ignored=ignored, bank=bank)
    err = (lambda *a: None) if quiet else (lambda *a: print(*a, file=sys.stderr))
    for f in res['flags']:
        if f['type'] == 'unknown_department_code':
            err(f"WARNING: unknown department code {f['department_code']} — placeholder used, confirm before posting")
        else:
            err(f"WARNING: unmapped label '{f['label']}' in dept {f['department_code']} ({f['amount']}) — flagged {UNMAPPED_MARK}")
    if res['failures']:
        err('VALIDATION FAILED — output withheld, human review required:')
        for line in res['failures']:
            err('  ' + line)
        if not force:
            return res, None
        err('  --force given: writing files anyway for review.')
    import os
    ref = os.path.join(out_dir, f'PYROLL_REFERENCE_{header_date}.xlsx')
    draft = os.path.join(out_dir, f'{header_date}_by_department_DRAFT.xlsx')
    write_reference(res, ref)
    write_draft(res, draft)
    if not quiet:
        print(f'Wrote {ref}')
        print(f'Wrote {draft}')
        print(f"Department matrix total = closing total = {res['matrix_total']}")
    return res, (ref, draft)


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument('csv_path')
    p.add_argument('--out-dir', default='.')
    p.add_argument('--bank-leg-a', type=str, help='bank withdrawal for Dues/GST/PST')
    p.add_argument('--bank-leg-b', type=str, help='bank withdrawal for EI/CPP/FedTax')
    p.add_argument('--bank-net-pay', type=str, help='bank withdrawal for Net Pay')
    p.add_argument('--force', action='store_true', help='write files even if validation fails')
    a = p.parse_args(argv)
    bank = None
    if a.bank_leg_a or a.bank_leg_b or a.bank_net_pay:
        bank = {'leg_a': a.bank_leg_a, 'leg_b': a.bank_leg_b, 'net_pay': a.bank_net_pay}
    res, files = run(a.csv_path, a.out_dir, bank=bank, force=a.force)
    if res['failures']:
        return 2
    if res['flags']:
        print(f"{len(res['flags'])} item(s) need review — see NOTES section in the DRAFT workbook.", file=sys.stderr)
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
