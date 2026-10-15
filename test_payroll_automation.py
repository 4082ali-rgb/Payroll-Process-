#!/usr/bin/env python3
"""Tests for payroll_automation.py against test_data/expected_results.json.
Run: python -m unittest test_payroll_automation -v
"""
import csv
import json
import os
import subprocess
import sys
import tempfile
import unittest
from collections import OrderedDict
from decimal import Decimal

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, 'test_data')
sys.path.insert(0, HERE)
import payroll_automation as pa  # noqa: E402
import payroll_je  # noqa: E402

with open(os.path.join(DATA, 'expected_results.json')) as f:
    EXPECTED = json.load(f)


def D(x):
    return Decimal(str(x)).quantize(Decimal('0.01'))


class BalancedCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.exp = EXPECTED['test_case_1_balanced']
        cls.tmp = tempfile.mkdtemp()
        cls.res, cls.files = pa.run(os.path.join(DATA, 'test_case_1_balanced.csv'), cls.tmp, quiet=True)

    def test_header_and_no_failures(self):
        self.assertEqual(self.res['header_date'], self.exp['pay_period_header'])
        self.assertEqual(self.res['failures'], [])
        self.assertEqual(self.res['flags'], [])
        self.assertIsNotNone(self.files)

    def test_department_matrix_every_figure(self):
        self.assertEqual(list(self.res['matrix']), list(self.exp['department_matrix']))
        for code, cats in self.exp['department_matrix'].items():
            for acct, (json_key, _qb) in pa.CATEGORY_META.items():
                self.assertEqual(self.res['matrix'][code][acct], D(cats[json_key]), (code, json_key))
        self.assertEqual(self.res['matrix_total'], D(self.exp['department_matrix_grand_total']))

    def test_closing_entries(self):
        for key, val in self.exp['closing_entries'].items():
            self.assertEqual(self.res['closing'][key], D(val), key)
        self.assertEqual(self.res['closing_total'], D(self.exp['closing_entries_total']))
        self.assertEqual(self.res['closing_total'], self.res['matrix_total'])

    def test_legs(self):
        for key, val in self.exp['leg_a_dues_gst_pst'].items():
            self.assertEqual(self.res['leg_a'][key], D(val), key)
        for key, val in self.exp['leg_b_ei_cpp_fedtax'].items():
            self.assertEqual(self.res['leg_b'][key], D(val), key)

    def test_cross_checks(self):
        cc = self.exp['cross_checks']
        self.assertEqual(self.res['category_totals'][6046],
                         D(cc['cpp_expense_matrix_total_must_equal_top_block_cpp_employer_total']['matrix_side']))
        self.assertEqual(self.res['category_totals'][6046], self.res['cpp_employer_top'])
        self.assertEqual(self.res['category_totals'][6047],
                         D(cc['ei_expense_matrix_total_must_equal_top_block_ei_employer']['matrix_side']))
        self.assertEqual(self.res['ei_employer_top'],
                         D(cc['ei_expense_matrix_total_must_equal_top_block_ei_employer']['top_block_side']))

    def test_health_benefits_zero_is_noted_not_flagged(self):
        self.assertEqual(self.res['closing']['6043_employer_health_benefits'], Decimal('0'))
        self.assertTrue(any('confirmed $0' in n for n in self.res['notes']))

    def test_workbooks_written_with_formulas(self):
        from openpyxl import load_workbook
        ref, draft = self.files
        self.assertTrue(os.path.basename(ref) == 'PYROLL_REFERENCE_20260901.xlsx')
        self.assertTrue(os.path.basename(draft) == '20260901_by_department_DRAFT.xlsx')
        ws = load_workbook(ref).active
        self.assertEqual(str(ws['A1'].value), '20260901')
        formulas = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('=')]
        self.assertGreater(len(formulas), 30)
        ws2 = load_workbook(draft).active
        self.assertIn('DRAFT — built from raw payroll CSV, not a posted JE.', ws2['A2'].value)

    def test_end_to_end_draft_to_qb_csv(self):
        """DRAFT workbook -> payroll_je.py must produce a balanced, fully-mapped CSV."""
        _ref, draft = self.files
        rows, warnings, flagged, dr, cr = payroll_je.convert(draft, 'JJ0001', '5/9/2026')
        self.assertEqual(warnings, [])
        self.assertEqual(flagged, 0)
        self.assertEqual(dr, cr)
        self.assertEqual(dr, D(self.exp['department_matrix_grand_total']))
        classes = {r[9] for r in rows if r[4]}
        self.assertEqual(classes, {'0010-ACCOMODATIONS', '0015-HOUSEKEEPING'})
        self.assertTrue(all(r[9] == '' for r in rows if r[5]))

    def test_cli_exit_code(self):
        p = subprocess.run([sys.executable, os.path.join(HERE, 'payroll_automation.py'),
                            os.path.join(DATA, 'test_case_1_balanced.csv'), '--out-dir', self.tmp],
                           capture_output=True, text=True)
        self.assertEqual(p.returncode, 0, p.stderr)

    def test_bank_mismatch_blocks_output(self):
        tmp = tempfile.mkdtemp()
        res, files = pa.run(os.path.join(DATA, 'test_case_1_balanced.csv'), tmp,
                            bank={'leg_a': '525.50', 'leg_b': '12160.00', 'net_pay': '14700.00'}, quiet=True)
        self.assertIsNone(files)
        self.assertTrue(any('CHECK 6' in f for f in res['failures']))
        res, files = pa.run(os.path.join(DATA, 'test_case_1_balanced.csv'), tmp,
                            bank={'leg_a': '525.50', 'leg_b': '12160.00', 'net_pay': '14500.00'}, quiet=True)
        self.assertIsNotNone(files)
        self.assertTrue(res['bank_confirmed'])


class UnknownsCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.exp = EXPECTED['test_case_2_unknowns']['expected_flags']
        cls.tmp = tempfile.mkdtemp()
        cls.res, cls.files = pa.run(os.path.join(DATA, 'test_case_2_unknowns.csv'), cls.tmp, force=True, quiet=True)

    def test_flags_match_expected(self):
        got = {(f['type'], f['department_code'], f.get('label'), f.get('amount')) for f in self.res['flags']}
        want = {(e['type'], e['department_code'], e.get('label'),
                 D(e['amount']) if 'amount' in e else None) for e in self.exp}
        self.assertEqual(got, want)

    def test_placeholders_never_guess(self):
        self.assertEqual(self.res['dept_names']['99'], 'DEPT 99 (name not confirmed)')
        self.assertEqual(self.res['unmapped'], [('99', 'Mystery Allowance', Decimal('-500.00'))])
        # not folded into wages or anything else
        self.assertEqual(sum(self.res['matrix']['99'].values()), Decimal('0'))

    def test_validation_blocks_without_force(self):
        res, files = pa.run(os.path.join(DATA, 'test_case_2_unknowns.csv'), tempfile.mkdtemp(), quiet=True)
        self.assertIsNone(files)
        self.assertTrue(res['failures'])

    def test_draft_shows_placeholders_highlighted(self):
        from openpyxl import load_workbook
        ws = load_workbook(self.files[1]).active
        cells = [c for row in ws.iter_rows() for c in row if c.value is not None]
        hdr = next(c for c in cells if c.value == '99 — DEPT 99 (name not confirmed)')
        self.assertEqual(hdr.fill.fgColor.rgb[-6:], 'FFC7CE')
        unm = next(c for c in cells if c.value == pa.UNMAPPED_MARK)
        self.assertEqual(unm.fill.fgColor.rgb[-6:], 'FFC7CE')
        notes = [c.value for c in cells if isinstance(c.value, str) and 'Mystery Allowance' in c.value and 'NEEDS' not in c.value]
        self.assertTrue(any("department 99" in n and '-500.00' in n for n in notes))
        ws_ref = load_workbook(self.files[0]).active
        self.assertTrue(any(c.value == '99 — DEPT 99 (name not confirmed)' and c.fill.fgColor.rgb[-6:] == 'FFC7CE'
                            for row in ws_ref.iter_rows() for c in row))


class UnmappedTopBlockLabel(unittest.TestCase):
    """A top-block (company-wide) label that isn't in Leg A, Leg B, or any closing
    formula must be flagged, never silently dropped from the output."""

    def test_unknown_topblock_label_is_flagged(self):
        top_block = [
            ('B10657', 'Net Pay', Decimal('1000.00')),
            ('B10657', 'Service Fees', Decimal('10.00')),
            ('B10657', 'GST', Decimal('0.50')),
            ('B10657', 'PST', Decimal('0.10')),
            ('B10657', 'Gym Membership Reimbursement', Decimal('-500.00')),
        ]
        res = pa.compute('20260101', top_block, {}, [], ignored=[])
        self.assertEqual(res['unmapped_top'], [('Gym Membership Reimbursement', Decimal('-500.00'))])
        self.assertIn({'type': 'unmapped_topblock_label', 'label': 'Gym Membership Reimbursement',
                       'amount': Decimal('-500.00')}, res['flags'])
        lines = pa.review_lines(res)
        self.assertTrue(any('Gym Membership Reimbursement' in l and 'NOT included' in l for l in lines))

    def test_pet_deposit_maps_to_staff_deposits_and_balances(self):
        """Pet Deposit (confirmed 2026-10) folds into 2024 Staff Deposits like the
        existing company-wide 'Staff Deposit' line. A negative net figure there must
        land in the Debit column of the DRAFT workbook, never as a negative Credit."""
        top_block = [
            ('B10657', 'Net Pay', Decimal('1000.00')),
            ('B10657', 'Pet Deposit', Decimal('-500.00')),
        ]
        res = pa.compute('20260101', top_block, OrderedDict(), [], ignored=[])
        self.assertEqual(res['unmapped_top'], [])
        self.assertEqual(res['closing']['2024_staff_deposit_company'], Decimal('-500.00'))
        tmp = tempfile.mkdtemp()
        draft = os.path.join(tmp, 'draft.xlsx')
        pa.write_draft(res, draft)
        from openpyxl import load_workbook
        ws = load_workbook(draft).active
        row = next(r for r in ws.iter_rows() if r[0].value == '2024 Staff Deposits')
        self.assertEqual(row[1].value, 500.0)   # Debit column, positive
        self.assertIsNone(row[2].value)         # Credit column empty

    def test_known_topblock_labels_never_flagged(self):
        """Regression guard: every label actually consumed by Leg A/B/closing must
        NOT be flagged as unmapped (this is what the existing fixtures exercise)."""
        res, _files = pa.run(os.path.join(DATA, 'test_case_1_balanced.csv'), tempfile.mkdtemp(), quiet=True)
        self.assertEqual(res.get('unmapped_top', []), [])


class LegJournalCsv(unittest.TestCase):
    """leg_journal_csv.py: the standalone Leg A / Leg B QuickBooks JE exports."""

    @classmethod
    def setUpClass(cls):
        import leg_journal_csv as lj
        cls.lj = lj
        cls.exp = EXPECTED['test_case_1_balanced']
        cls.tmp = tempfile.mkdtemp()

    def _rows(self, leg):
        out = os.path.join(self.tmp, f'leg{leg}.csv')
        rc = self.lj.main([os.path.join(DATA, 'test_case_1_balanced.csv'), '--leg', leg,
                           '--journal-no', 'JJTEST', '--date', '1/9/2026', '--output', out])
        self.assertEqual(rc, 0)
        with open(out, newline='') as f:
            return list(csv.DictReader(f))

    def test_leg_a_matches_expected_figures(self):
        rows = self._rows('a')
        self.assertEqual(len(rows), 4)
        exp = self.exp['leg_a_dues_gst_pst']
        got = {r['*AccountName']: Decimal(r['Debits'] or f"-{r['Credits']}") for r in rows}
        self.assertEqual(got['5004 Dues and Subscriptions'], Decimal(str(exp['5004_dues_and_subscription'])))
        self.assertEqual(got['2031 GST Paid on Purchases'], Decimal(str(exp['2031_gst_paid_on_purchases'])))
        self.assertEqual(got['2040 PST paid on Purchases'], Decimal(str(exp['2040_pst_paid_on_purchases'])))
        self.assertEqual(-got['1004 Envision Bank Account'], Decimal(str(exp['1004_envision_bank_account_credit'])))
        self.assertTrue(all(r['Class'] == '0090-GENERAL & ADMIN' for r in rows if r['Debits']))
        self.assertTrue(all(r['Class'] == '' for r in rows if r['Credits']))

    def test_leg_b_matches_expected_figures_and_balances(self):
        rows = self._rows('b')
        self.assertEqual(len(rows), 4)
        exp = self.exp['leg_b_ei_cpp_fedtax']
        got = {r['*AccountName']: Decimal(r['Debits'] or f"-{r['Credits']}") for r in rows}
        self.assertEqual(got['2013 EI Payable'], Decimal(str(exp['2013_ei_payable'])))
        self.assertEqual(got['2014 CPP Payable'], Decimal(str(exp['2014_cpp_payable'])))
        self.assertEqual(got['2017 Federal Income Tax Payable'], Decimal(str(exp['2017_federal_income_tax_payable'])))
        self.assertEqual(-got['1004 Envision Bank Account'], Decimal(str(exp['1004_envision_bank_account_credit'])))
        debits = sum(Decimal(r['Debits']) for r in rows if r['Debits'])
        credits = sum(Decimal(r['Credits']) for r in rows if r['Credits'])
        self.assertEqual(debits, credits)

    def test_custom_class_override(self):
        out = os.path.join(self.tmp, 'legA_noclass.csv')
        self.lj.main([os.path.join(DATA, 'test_case_1_balanced.csv'), '--leg', 'a',
                     '--journal-no', 'JJTEST', '--date', '1/9/2026', '--output', out, '--class', ''])
        with open(out, newline='') as f:
            rows = list(csv.DictReader(f))
        self.assertTrue(all(r['Class'] == '' for r in rows))

    def test_memo_defaults_distinguish_the_legs(self):
        rows_a, rows_b = self._rows('a'), self._rows('b')
        self.assertTrue(all(r['Memo'] == 'Dues and Subscriptions' for r in rows_a))
        self.assertTrue(all(r['Memo'] == 'EI, CPP & Federal Tax' for r in rows_b))
        self.assertNotEqual(rows_a[0]['Memo'], rows_b[0]['Memo'])

    def test_memo_override(self):
        out = os.path.join(self.tmp, 'legA_memo.csv')
        self.lj.main([os.path.join(DATA, 'test_case_1_balanced.csv'), '--leg', 'a',
                     '--journal-no', 'JJTEST', '--date', '1/9/2026', '--output', out, '--memo', 'Custom Note'])
        with open(out, newline='') as f:
            rows = list(csv.DictReader(f))
        self.assertTrue(all(r['Memo'] == 'Custom Note' for r in rows))


class PayrollJeMemo(unittest.TestCase):
    def test_memo_blank_by_default_and_settable(self):
        fixture = os.path.join(HERE, 'test_fixture.xlsx')
        rows, _w, _f, _dr, _cr = payroll_je.convert(fixture, 'TEST001', '1/1/2026')
        self.assertTrue(all(r[2] == '' for r in rows))
        rows, _w, _f, _dr, _cr = payroll_je.convert(fixture, 'TEST001', '1/1/2026', memo='Department Breakdown')
        self.assertTrue(all(r[2] == 'Department Breakdown' for r in rows))


if __name__ == '__main__':
    unittest.main()
